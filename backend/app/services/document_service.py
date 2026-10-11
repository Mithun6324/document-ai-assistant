import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional
from uuid import UUID, uuid4

import pdfplumber
from docx import Document as DocxDocument
from fastapi import BackgroundTasks, HTTPException, UploadFile, status
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import AuditLog, Chunk, Document, Message
from app.schemas import DocumentQuestionItem, DocumentResponse, DocumentStatusResponse
from app.services.chroma_service import get_chroma_service
from app.services.chunk_service import semantic_chunk
from app.services.cleaning_service import clean_text
from app.services.embedding_service import get_embedding_service
from app.utils.file_validation import estimate_tokens, file_sha256, validate_upload
from app.utils.logging import get_logger

logger = get_logger(__name__)
settings = get_settings()


def _extract_pdf(path: Path) -> list[dict]:
    pages = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            pages.append({"page_number": i, "text": clean_text(text)})
    return pages


def _extract_docx(path: Path) -> list[dict]:
    doc = DocxDocument(path)
    paragraphs = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            paragraphs.append(text)
    full = clean_text("\n\n".join(paragraphs))
    return [{"page_number": 1, "text": full}] if full else []


def _update_progress(db: Session, doc: Document, stage: str, percent: int) -> None:
    doc.progress_stage = stage
    doc.progress_percent = percent
    doc.status = "processing"
    doc.updated_at = datetime.now(timezone.utc)
    db.commit()


def process_document(document_id: str) -> None:
    """Background processing pipeline."""
    from app.db.session import SessionLocal

    db = SessionLocal()
    try:
        doc = db.query(Document).filter(Document.id == UUID(document_id)).first()
        if not doc or doc.deleted_at:
            return

        path = Path(doc.file_path)
        _update_progress(db, doc, "extracting", 10)

        if doc.file_type == "pdf":
            pages = _extract_pdf(path)
        else:
            pages = _extract_docx(path)

        if not any(p["text"].strip() for p in pages):
            doc.status = "failed"
            doc.processing_error = "Empty document: no extractable text found"
            doc.progress_stage = "failed"
            doc.progress_percent = 100
            db.commit()
            return

        doc.page_count = len(pages)
        _update_progress(db, doc, "cleaning", 30)

        # Cleaning already applied during extract
        _update_progress(db, doc, "chunking", 50)
        chunks = semantic_chunk(pages, str(doc.id), doc.original_filename)
        if not chunks:
            doc.status = "failed"
            doc.processing_error = "No chunks produced from document"
            doc.progress_stage = "failed"
            doc.progress_percent = 100
            db.commit()
            return

        _update_progress(db, doc, "embedding", 70)
        texts = [c["text"] for c in chunks]
        embeddings = get_embedding_service().embed_texts(texts)

        _update_progress(db, doc, "indexing", 85)

        # Persist chunks in Postgres
        db.query(Chunk).filter(Chunk.document_id == doc.id).delete()
        for c in chunks:
            db.add(
                Chunk(
                    chunk_id=c["chunk_id"],
                    document_id=doc.id,
                    text=c["text"],
                    page_number=c.get("page_number"),
                    section_name=c.get("section_name"),
                    chunk_type=c.get("chunk_type", "paragraph"),
                    token_estimate=estimate_tokens(c["text"]),
                )
            )

        metadatas = []
        for c in chunks:
            metadatas.append(
                {
                    "chunk_id": c["chunk_id"],
                    "document_id": str(doc.id),
                    "document_name": doc.original_filename,
                    "user_id": str(doc.user_id),
                    "page_number": c.get("page_number") or 0,
                    "section_name": c.get("section_name") or "",
                    "chunk_type": c.get("chunk_type") or "paragraph",
                    "created_at": c.get("created_at") or "",
                }
            )

        get_chroma_service().add_chunks(
            chunk_ids=[c["chunk_id"] for c in chunks],
            texts=texts,
            embeddings=embeddings,
            metadatas=metadatas,
        )

        doc.chunk_count = len(chunks)
        doc.status = "ready"
        doc.progress_stage = "ready"
        doc.progress_percent = 100
        doc.processing_error = None
        doc.updated_at = datetime.now(timezone.utc)
        db.add(
            AuditLog(
                user_id=doc.user_id,
                action="document_processed",
                metadata_json={"document_id": str(doc.id), "chunks": len(chunks)},
            )
        )
        db.commit()
        logger.info("Document %s ready with %s chunks", doc.id, len(chunks))
    except Exception as exc:
        logger.exception("Document processing failed: %s", document_id)
        try:
            doc = db.query(Document).filter(Document.id == UUID(document_id)).first()
            if doc:
                doc.status = "failed"
                doc.progress_stage = "failed"
                doc.progress_percent = 100
                doc.processing_error = str(exc)
                db.commit()
        except Exception:
            pass
    finally:
        db.close()


def _question_count(db: Session, document_id: UUID) -> int:
    # Count user messages whose retrieved_chunks reference this document
    messages = (
        db.query(Message)
        .filter(Message.role == "user", Message.retrieved_chunks.isnot(None))
        .all()
    )
    count = 0
    doc_str = str(document_id)
    for m in messages:
        chunks = m.retrieved_chunks or []
        if any(str(c.get("document_id")) == doc_str for c in chunks if isinstance(c, dict)):
            count += 1
    return count


def _to_response(db: Session, doc: Document, duplicate: bool = False) -> DocumentResponse:
    return DocumentResponse(
        id=doc.id,
        filename=doc.filename,
        original_filename=doc.original_filename,
        file_type=doc.file_type,
        file_size=doc.file_size,
        page_count=doc.page_count,
        chunk_count=doc.chunk_count,
        status=doc.status,
        progress_stage=doc.progress_stage,
        progress_percent=doc.progress_percent,
        processing_error=doc.processing_error,
        question_count=_question_count(db, doc.id),
        created_at=doc.created_at,
        updated_at=doc.updated_at,
        duplicate=duplicate,
    )


async def upload_document(
    db: Session,
    user_id: UUID,
    file: UploadFile,
    background_tasks: BackgroundTasks,
) -> DocumentResponse:
    content = await file.read()
    file_type = validate_upload(file, content)
    file_hash = file_sha256(content)

    existing = (
        db.query(Document)
        .filter(
            Document.user_id == user_id,
            Document.file_hash == file_hash,
            Document.deleted_at.is_(None),
            Document.status == "ready",
        )
        .first()
    )
    if existing:
        logger.info("Duplicate upload detected for user %s hash %s", user_id, file_hash[:12])
        return _to_response(db, existing, duplicate=True)

    upload_root = Path(settings.upload_dir) / str(user_id)
    upload_root.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid4()}.{file_type}"
    dest = upload_root / stored_name
    dest.write_bytes(content)

    doc = Document(
        user_id=user_id,
        filename=stored_name,
        original_filename=file.filename or stored_name,
        file_type=file_type,
        file_hash=file_hash,
        file_path=str(dest),
        file_size=len(content),
        status="pending",
        progress_stage="queued",
        progress_percent=0,
    )
    db.add(doc)
    db.add(
        AuditLog(
            user_id=user_id,
            action="document_upload",
            metadata_json={"filename": doc.original_filename, "size": len(content)},
        )
    )
    db.commit()
    db.refresh(doc)

    background_tasks.add_task(process_document, str(doc.id))
    logger.info("Queued document %s for processing", doc.id)
    return _to_response(db, doc)


def list_documents(db: Session, user_id: UUID) -> list[DocumentResponse]:
    docs = (
        db.query(Document)
        .filter(Document.user_id == user_id, Document.deleted_at.is_(None))
        .order_by(Document.created_at.desc())
        .all()
    )
    return [_to_response(db, d) for d in docs]


def get_document(db: Session, user_id: UUID, document_id: UUID) -> DocumentResponse:
    doc = (
        db.query(Document)
        .filter(Document.id == document_id, Document.user_id == user_id, Document.deleted_at.is_(None))
        .first()
    )
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return _to_response(db, doc)


def get_document_status(db: Session, user_id: UUID, document_id: UUID) -> DocumentStatusResponse:
    doc = (
        db.query(Document)
        .filter(Document.id == document_id, Document.user_id == user_id, Document.deleted_at.is_(None))
        .first()
    )
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
    return DocumentStatusResponse(
        id=doc.id,
        status=doc.status,
        progress_stage=doc.progress_stage,
        progress_percent=doc.progress_percent,
        processing_error=doc.processing_error,
        chunk_count=doc.chunk_count,
        page_count=doc.page_count,
    )


def soft_delete_document(db: Session, user_id: UUID, document_id: UUID) -> dict:
    doc = (
        db.query(Document)
        .filter(Document.id == document_id, Document.user_id == user_id, Document.deleted_at.is_(None))
        .first()
    )
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    get_chroma_service().delete_by_document(str(doc.id), str(user_id))
    doc.deleted_at = datetime.now(timezone.utc)
    doc.status = "deleted"
    db.add(AuditLog(user_id=user_id, action="document_delete", metadata_json={"document_id": str(doc.id)}))
    db.commit()
    return {"detail": "Document deleted"}


def get_document_questions(db: Session, user_id: UUID, document_id: UUID) -> list[DocumentQuestionItem]:
    doc = (
        db.query(Document)
        .filter(Document.id == document_id, Document.user_id == user_id, Document.deleted_at.is_(None))
        .first()
    )
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    from app.models import Conversation

    conversations = db.query(Conversation).filter(Conversation.user_id == user_id).all()
    items: list[DocumentQuestionItem] = []
    doc_str = str(document_id)

    for conv in conversations:
        msgs = sorted(conv.messages, key=lambda m: m.created_at)
        for i, m in enumerate(msgs):
            if m.role != "assistant":
                continue
            sources = m.sources or []
            retrieved = m.retrieved_chunks or []
            linked = any(str(s.get("document_id")) == doc_str for s in sources if isinstance(s, dict))
            linked = linked or any(
                str(c.get("metadata", {}).get("document_id") if isinstance(c.get("metadata"), dict) else c.get("document_id"))
                == doc_str
                for c in retrieved
                if isinstance(c, dict)
            )
            if not linked:
                continue
            q_msg = msgs[i - 1] if i > 0 and msgs[i - 1].role == "user" else None
            items.append(
                DocumentQuestionItem(
                    message_id=(q_msg.id if q_msg else m.id),
                    conversation_id=conv.id,
                    question=(q_msg.content if q_msg else "(unknown)"),
                    answer_preview=m.content[:200],
                    asked_at=(q_msg.created_at if q_msg else m.created_at),
                )
            )

    items.sort(key=lambda x: x.asked_at, reverse=True)
    return items


def suggested_questions(db: Session, user_id: UUID, document_id: UUID) -> list[str]:
    doc = (
        db.query(Document)
        .filter(Document.id == document_id, Document.user_id == user_id, Document.deleted_at.is_(None))
        .first()
    )
    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    sample = (
        db.query(Chunk)
        .filter(Chunk.document_id == doc.id)
        .limit(8)
        .all()
    )
    joined = " ".join(c.section_name or "" for c in sample) + " " + " ".join(c.text[:200] for c in sample)
    lower = joined.lower()

    if any(k in lower for k in ["skills", "experience", "education", "certification", "resume"]):
        return [
            "Summarize this resume",
            "List the key skills",
            "Summarize work experience",
            "List certifications and education",
        ]
    if any(k in lower for k in ["revenue", "balance sheet", "income", "cash flow", "financial"]):
        return [
            "Summarize the key financial figures",
            "What is the total revenue?",
            "List major expenses",
            "Highlight important risks or notes",
        ]
    return [
        "Summarize this document",
        "What are the main topics?",
        "List key points",
        "Extract important dates and names",
    ]
